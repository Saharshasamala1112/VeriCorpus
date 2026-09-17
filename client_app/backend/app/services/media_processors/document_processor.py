from __future__ import annotations

import hashlib
from pathlib import Path

from app.core.media_registry import get_capability
from app.services.media_processors import (
    MediaProcessor,
    ProcessingContext,
    ProcessingResult,
    register_processor,
)


class DocumentProcessor(MediaProcessor):
    modality = "document"

    async def validate(self, ctx: ProcessingContext) -> ProcessingResult:
        path = Path(ctx.file_path)
        if not path.exists():
            return ProcessingResult(success=False, errors=["File not found"])

        if ctx.file_size == 0:
            return ProcessingResult(success=False, errors=["Empty file"])

        capability = get_capability("document")
        max_size = capability.max_size_bytes
        if ctx.file_size > max_size:
            return ProcessingResult(
                success=False,
                errors=[f"Document too large: {ctx.file_size} bytes (max {max_size})"],
            )

        ext = path.suffix.lower()
        if ext and not capability.accepts(path.name, ctx.mime_type):
            return ProcessingResult(
                success=False,
                errors=[f"Unsupported document format: {ext}"],
            )

        return ProcessingResult(success=True)

    async def extract_metadata(self, ctx: ProcessingContext) -> ProcessingResult:
        path = Path(ctx.file_path)
        ext = path.suffix.lower()

        try:
            if ext == ".pdf":
                return await self._extract_pdf_metadata(ctx)
            elif ext == ".docx":
                return await self._extract_docx_metadata(ctx)
            else:
                return await self._extract_text_metadata(ctx)
        except Exception as e:
            return ProcessingResult(success=False, errors=[f"Failed to extract document metadata: {e}"])

    async def _extract_pdf_metadata(self, ctx: ProcessingContext) -> ProcessingResult:
        path = Path(ctx.file_path)
        try:
            import fitz  # PyMuPDF

            doc = fitz.open(str(path))
            metadata = {
                "page_count": doc.page_count,
                "format": "pdf",
                "title": doc.metadata.get("title", ""),
                "author": doc.metadata.get("author", ""),
                "subject": doc.metadata.get("subject", ""),
                "creator": doc.metadata.get("creator", ""),
                "producer": doc.metadata.get("producer", ""),
                "encrypted": doc.is_encrypted,
                "text_length": sum(len(page.get_text()) for page in doc),
            }
            doc.close()
            ctx.metadata.update(metadata)
            return ProcessingResult(success=True, metadata=metadata)

        except ImportError:
            return ProcessingResult(
                success=False,
                errors=["PyMuPDF not installed. Install with: pip install PyMuPDF"],
            )

    async def _extract_docx_metadata(self, ctx: ProcessingContext) -> ProcessingResult:
        path = Path(ctx.file_path)
        try:
            from docx import Document

            doc = Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs]
            full_text = "\n".join(paragraphs)

            core_props = doc.core_properties
            metadata = {
                "format": "docx",
                "title": core_props.title or "",
                "author": core_props.author or "",
                "subject": core_props.subject or "",
                "paragraph_count": len(doc.paragraphs),
                "table_count": len(doc.tables),
                "image_count": len(doc.inline_shapes),
                "text_length": len(full_text),
            }
            ctx.metadata.update(metadata)
            return ProcessingResult(success=True, metadata=metadata)

        except ImportError:
            return ProcessingResult(
                success=False,
                errors=["python-docx not installed. Install with: pip install python-docx"],
            )

    async def _extract_text_metadata(self, ctx: ProcessingContext) -> ProcessingResult:
        path = Path(ctx.file_path)
        try:
            text = path.read_text(encoding="utf-8", errors="replace")
            lines = text.split("\n")
            words = text.split()
            metadata = {
                "format": path.suffix.lstrip("."),
                "line_count": len(lines),
                "word_count": len(words),
                "text_length": len(text),
            }
            ctx.metadata.update(metadata)
            return ProcessingResult(success=True, metadata=metadata)

        except Exception as e:
            return ProcessingResult(success=False, errors=[f"Failed to read text document: {e}"])

    async def normalize(self, ctx: ProcessingContext) -> ProcessingResult:
        path = Path(ctx.file_path)
        ext = path.suffix.lower()

        if ext == ".pdf":
            return await self._normalize_pdf(ctx)
        elif ext == ".docx":
            return await self._normalize_docx(ctx)
        else:
            ctx.metadata["normalized_path"] = str(path)
            return ProcessingResult(success=True, normalized_path=str(path))

    async def _normalize_pdf(self, ctx: ProcessingContext) -> ProcessingResult:
        path = Path(ctx.file_path)
        try:
            import fitz

            doc = fitz.open(str(path))
            text_content = []
            for page in doc:
                text_content.append(page.get_text())
            doc.close()

            full_text = "\n\n".join(text_content)
            normalized_path = ctx.file_path + ".normalized.txt"
            Path(normalized_path).write_text(full_text, encoding="utf-8")
            ctx.metadata["normalized_path"] = normalized_path
            return ProcessingResult(success=True, normalized_path=normalized_path)

        except ImportError:
            return ProcessingResult(
                success=False,
                errors=["PyMuPDF not installed. Install with: pip install PyMuPDF"],
            )

    async def _normalize_docx(self, ctx: ProcessingContext) -> ProcessingResult:
        path = Path(ctx.file_path)
        try:
            from docx import Document

            doc = Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs]
            full_text = "\n".join(paragraphs)

            normalized_path = ctx.file_path + ".normalized.txt"
            Path(normalized_path).write_text(full_text, encoding="utf-8")
            ctx.metadata["normalized_path"] = normalized_path
            return ProcessingResult(success=True, normalized_path=normalized_path)

        except ImportError:
            return ProcessingResult(
                success=False,
                errors=["python-docx not installed. Install with: pip install python-docx"],
            )

    async def preprocess(self, ctx: ProcessingContext) -> ProcessingResult:
        normalized_path = ctx.metadata.get("normalized_path", ctx.file_path)
        try:
            text = Path(normalized_path).read_text(encoding="utf-8", errors="replace")

            import re

            text = re.sub(r"\n{3,}", "\n\n", text)
            text = re.sub(r"[ \t]+", " ", text)
            text = text.strip()

            preprocessed_path = ctx.file_path + ".preprocessed.txt"
            Path(preprocessed_path).write_text(text, encoding="utf-8")
            ctx.metadata["preprocessed_path"] = preprocessed_path
            ctx.metadata["preprocessed_length"] = len(text)
            return ProcessingResult(success=True, normalized_path=preprocessed_path)

        except Exception as e:
            return ProcessingResult(success=False, errors=[f"Failed to preprocess document: {e}"])

    async def extract_features(self, ctx: ProcessingContext) -> ProcessingResult:
        preprocessed_path = ctx.metadata.get("preprocessed_path", ctx.metadata.get("normalized_path", ctx.file_path))
        try:
            text = Path(preprocessed_path).read_text(encoding="utf-8", errors="replace")
            words = text.split()
            sentences = [s.strip() for s in text.split(".") if s.strip()]
            paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]

            features = {
                "word_count": len(words),
                "sentence_count": len(sentences),
                "paragraph_count": len(paragraphs),
                "avg_word_length": round(sum(len(w) for w in words) / max(len(words), 1), 2),
                "text_length": len(text),
                "file_hash": hashlib.sha256(Path(ctx.file_path).read_bytes()).hexdigest(),
            }
            ctx.features.update(features)
            return ProcessingResult(success=True, features=features)

        except Exception as e:
            return ProcessingResult(success=False, errors=[f"Failed to extract document features: {e}"])

    async def parse_structure(self, ctx: ProcessingContext) -> ProcessingResult:
        path = Path(ctx.file_path)
        structure: dict[str, object] = {"pages": [], "paragraphs": [], "headings": [], "tables": []}
        try:
            if path.suffix.lower() == ".pdf":
                import fitz

                doc = fitz.open(str(path))
                for page_number, page in enumerate(doc, start=1):
                    blocks = page.get_text("blocks")
                    structure["pages"].append({"page": page_number, "block_count": len(blocks)})
                    structure["paragraphs"].extend(
                        {"page": page_number, "text": block[4], "bbox": block[:4]}
                        for block in blocks
                        if block[4].strip()
                    )
                doc.close()
            elif path.suffix.lower() == ".docx":
                from docx import Document

                doc = Document(str(path))
                structure["paragraphs"] = [
                    {"text": p.text, "style": p.style.name} for p in doc.paragraphs if p.text.strip()
                ]
                structure["headings"] = [
                    {"text": p.text, "style": p.style.name}
                    for p in doc.paragraphs
                    if p.text.strip() and p.style.name.lower().startswith("heading")
                ]
                structure["tables"] = [
                    [[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables
                ]
            else:
                text = path.read_text(encoding="utf-8", errors="replace")
                structure["paragraphs"] = [{"text": paragraph} for paragraph in text.split("\n\n") if paragraph.strip()]
            ctx.metadata["structure"] = structure
            return ProcessingResult(success=True, metadata={"structure": structure})
        except (ImportError, OSError) as exc:
            return ProcessingResult(success=False, errors=[f"Failed to parse document structure: {exc}"])


register_processor("document", DocumentProcessor)
